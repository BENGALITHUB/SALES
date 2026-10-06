"""Builds branded, customisable Excel (.xlsx) lead reports."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

logger = config.get_logger(__name__)

# Report column label -> key inside a lead record.
FIELDS = {
    "Business Name": "name",
    "Address": "address",
    "City": "city",
    "Pincode": "pincode",
    "Phone Number": "phone",
    "Email ID": "email",
    "Website": "website",
    "Instagram": "instagram",
    "Facebook": "facebook",
}
COLUMNS = list(FIELDS)
CONTACT_COLUMNS = ("Phone Number", "Email ID", "Website", "Instagram", "Facebook")
URL_COLUMNS = ("Website", "Instagram", "Facebook")

# Theme name -> brand colour (hex, no #).
THEMES = {
    "Indigo": "4F46E5",
    "Emerald": "059669",
    "Ocean": "0369A1",
    "Crimson": "BE123C",
    "Graphite": "334155",
}

# Upper bound for each column's width; narrower content shrinks the column.
MAX_WIDTHS = {
    "#": 6,
    "Business Name": 34,
    "Address": 50,
    "City": 24,
    "Pincode": 12,
    "Phone Number": 20,
    "Email ID": 34,
    "Website": 38,
    "Instagram": 36,
    "Facebook": 36,
}

INK = "0F172A"
MUTED = "64748B"
GRID = Side(style="thin", color="E2E8F0")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)


def _tint(hex_color: str, amount: float) -> str:
    """Mix *hex_color* with white; amount=1 is pure white."""
    channels = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c + (255 - c) * amount):02X}" for c in channels)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def records_to_frame(records: list[dict]) -> pd.DataFrame:
    """Convert raw lead records into a DataFrame with report column labels."""
    rows = [{label: (r.get(key) or "") for label, key in FIELDS.items()} for r in records]
    return pd.DataFrame(rows, columns=COLUMNS)


def _banner(ws, text: str, row: int, last_col: str, *, font: Font, fill: PatternFill, height: float) -> None:
    ws.merge_cells(f"A{row}:{last_col}{row}")
    cell = ws[f"A{row}"]
    cell.value = text
    cell.font = font
    cell.fill = fill
    cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[row].height = height


def _write_leads_sheet(ws, df: pd.DataFrame, *, title: str, subtitle: str, brand: str, serial_numbers: bool) -> None:
    columns = (["#"] if serial_numbers else []) + list(df.columns)
    last_col = get_column_letter(len(columns))
    zebra = _fill(_tint(brand, 0.94))
    link_font = Font(color=brand, underline="single")

    ws.sheet_view.showGridLines = False
    _banner(ws, title, 1, last_col, font=Font(size=18, bold=True, color="FFFFFF"), fill=_fill(brand), height=38)
    _banner(ws, subtitle, 2, last_col, font=Font(size=10, italic=True, color="334155"), fill=_fill(_tint(brand, 0.86)), height=22)

    header_row = 4
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(header_row, col_idx, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(brand)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[header_row].height = 26

    for offset, values in enumerate(df.itertuples(index=False), start=1):
        row = header_row + offset
        values = ([offset] if serial_numbers else []) + list(values)
        for col_idx, (name, value) in enumerate(zip(columns, values), start=1):
            cell = ws.cell(row, col_idx, value if value != "" else None)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # never let scraped text become a formula
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="center",
                horizontal="center" if name == "#" else "left",
                wrap_text=name == "Address",
            )
            if offset % 2 == 0:
                cell.fill = zebra
            if not value:
                continue
            if name in URL_COLUMNS:
                cell.hyperlink = value
                cell.font = link_font
            elif name == "Email ID":
                cell.hyperlink = f"mailto:{value}"
                cell.font = link_font
            elif name == "Business Name":
                cell.font = Font(bold=True, color=INK)
            elif name == "#":
                cell.font = Font(color=MUTED)

    last_row = header_row + len(df)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    ws.freeze_panes = ws.cell(header_row + 1, 2 if serial_numbers else 1)

    for col_idx, name in enumerate(columns, start=1):
        cells = df[name] if name in df.columns else [str(len(df))]
        longest = max([len(name)] + [len(str(v)) for v in cells])
        width = max(min(longest + 3, MAX_WIDTHS.get(name, 40)), 8)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    note_row = last_row + 2
    ws.merge_cells(f"A{note_row}:{last_col}{note_row}")
    note = ws[f"A{note_row}"]
    note.value = (
        "Source: Google Places and publicly accessible business websites. "
        "Email and social links are best-effort and may not exist for every business."
    )
    note.font = Font(size=9, italic=True, color=MUTED)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.center.text = "Page &P of &N"


def _write_summary_sheet(ws, df: pd.DataFrame, *, title: str, brand: str, metadata: dict[str, Any]) -> None:
    light = _fill(_tint(brand, 0.92))
    section_font = Font(size=12, bold=True, color=brand)

    ws.sheet_view.showGridLines = False
    for letter, width in zip("ABC", (30, 34, 16)):
        ws.column_dimensions[letter].width = width
    _banner(ws, f"{title} — Summary", 1, "C", font=Font(size=16, bold=True, color="FFFFFF"), fill=_fill(brand), height=34)

    row = 3
    ws.cell(row, 1, "Search details").font = section_font
    details = {**{k: v for k, v in metadata.items() if v not in (None, "")}, "Leads in report": len(df)}
    for label, value in details.items():
        row += 1
        key_cell = ws.cell(row, 1, label)
        key_cell.font = Font(bold=True, color="334155")
        key_cell.fill = light
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        value_cell = ws.cell(row, 2, value)
        value_cell.alignment = Alignment(horizontal="left")
        for col in (1, 2, 3):
            ws.cell(row, col).border = BORDER

    row += 2
    ws.cell(row, 1, "Contact coverage").font = section_font
    row += 1
    for col_idx, name in enumerate(("Contact field", "Leads with data", "Coverage"), start=1):
        cell = ws.cell(row, col_idx, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(brand)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")

    first_data_row = row + 1
    total = len(df)
    for name in CONTACT_COLUMNS:
        if name not in df.columns:
            continue
        row += 1
        count = int((df[name].astype(str).str.strip() != "").sum())
        ws.cell(row, 1, name).font = Font(bold=True, color=INK)
        ws.cell(row, 2, count).alignment = Alignment(horizontal="center")
        pct = ws.cell(row, 3, count / total if total else 0)
        pct.number_format = "0%"
        pct.alignment = Alignment(horizontal="center")
        for col in (1, 2, 3):
            ws.cell(row, col).border = BORDER

    if row >= first_data_row:
        ws.conditional_formatting.add(
            f"C{first_data_row}:C{row}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=brand),
        )


def build_report(
    df: pd.DataFrame,
    *,
    title: str = "Lead Report",
    subtitle: str = "",
    theme: str = "Indigo",
    include_summary: bool = True,
    serial_numbers: bool = True,
    metadata: dict[str, Any] | None = None,
) -> bytes:
    """Render *df* (columns labelled like COLUMNS, any subset/order) into a
    styled workbook and return the .xlsx bytes."""
    brand = THEMES.get(theme, THEMES["Indigo"])
    df = df.fillna("").astype(str)
    subtitle = subtitle or f"Generated {datetime.now():%d %b %Y, %I:%M %p}"

    wb = Workbook()
    leads_ws = wb.active
    leads_ws.title = "Leads"
    _write_leads_sheet(leads_ws, df, title=title, subtitle=subtitle, brand=brand, serial_numbers=serial_numbers)
    if include_summary:
        _write_summary_sheet(wb.create_sheet("Summary"), df, title=title, brand=brand, metadata=metadata or {})

    wb.properties.title = title
    wb.properties.creator = "LeadScout"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_to_excel(records: list[dict], filename: str, **report_options: Any) -> str:
    """Write *records* to a styled .xlsx inside config.OUTPUT_DIR and return
    the full output path. *report_options* are passed to build_report."""
    output_path = config.OUTPUT_DIR / filename
    output_path.write_bytes(build_report(records_to_frame(records), **report_options))
    logger.info("Exported %d records to %s", len(records), output_path)
    return str(output_path)
